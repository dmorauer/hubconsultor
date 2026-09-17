import copy
import json
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engine import Project, file_hash

BASE = Path(__file__).resolve().parents[1]


def source_file(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    data = {
        'Empresas': [
            ['Nome*', 'CNPJ*', 'Nome para contato*', 'Telefone*', 'E-mail*', 'Moeda (BRL, EUR, USD)*', 'Código de integração', 'CEP*', 'Logradouro*', 'Número*', 'Bairro*', 'Cidade*', 'Estado*', 'País*'],
            ['Empresa Teste', '01.234.567/0001-95', 'Contato', '11 999999999', 'contato@example.com', 'BRL', 'UN1', '01.234-567', 'Rua Teste', '10', 'Centro', 'Cidade', 'SP', ' BRASIL'],
        ],
        'Colaboradores': [
            ['Nome completo*', 'Sexo (M ou F)', 'CPF*', 'E-mail*', 'Data de nascimento', 'Código de integração Casafruti', 'Código de integração Extrafruti', 'Ativo (S ou N)*', 'Usuário', 'Senha', 'Centro de custo (Cód. no ERP)', 'Descrição centro de custo', 'CPF Usuário aprovador', 'Nome Usuário aprovador', 'Cargo', 'Nome da Mãe', 'Telefone', 'RG', 'CNH', 'Data validade CNH', 'Passaporte', 'Data validade passaporte', 'Nacionalidade'],
            ['João da Silva', 'M', '01234567890', ' FINANCEIRO@example.com ', '01/02/1980', 100, 200, 'SIM'],
            ['João Silva', 'M', '11144477735', 'financeiro@example.com', '02/03/1981', None, 300, 'SIM'],
            ['Maria Silva', 'F', '52998224725', 'joao.silva@example.com', '03/04/1982', None, None, 'SIM'],
        ],
        'Centros de Custo': [
            ['Código Centro de custo*', 'Nome centro de Custo*', 'CPF Aprovador', 'Aprovador do centro de custo', 'Empresa(CNPJ)'],
            ['001', 'Administrativo', None, None, '01.234.567/0001-95'],
        ],
        'Tipos de Despesa': [
            ['Nome da Despesa*', 'Validação', 'Item de Orçamento', 'Conta Contábil', 'Valor Limite', 'Aplicação'],
            ['Hospedagem', 'Alerta', '0001', '123.4', '200', 'Por Dia'],
            ['OBS.: Estes registros são exemplos.', None, None, None, None, None],
        ],
    }
    for name, rows in data.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)
    wb.close()


class EngineTests(unittest.TestCase):
    def test_normalization_batch_is_explicit_stale_safe_and_persistent(self):
        self.p.original_values['USERS:2'][0] = '  João   Silva  '
        self.p.original_values['USERS:2'][4] = '2000-02-01'
        self.p.original_values['USERS:2'][8] = ' secret  '
        proposals = self.p.normalizations()
        self.assertEqual({p['category'] for p in proposals}, {'Espaços', 'Documentos', 'Datas', 'Sim/Não', 'CEP', 'Telefones'})
        self.assertFalse(any(p['kind'] == 'USERS' and p['col'] == 8 for p in proposals))
        self.assertFalse(self.p.decisions)
        self.p.accept_normalizations(proposals[:1])
        previous = dict(self.p.decisions)
        with self.assertRaises(ValueError):
            self.p.accept_normalizations(proposals)
        self.assertEqual(self.p.decisions, previous)
        self.p.accept_normalizations(self.p.normalizations())
        self.assertFalse(self.p.normalizations())
        self.assertEqual(self.p.effective()['USERS'][0].values[0], 'João Silva')
        session = self.root / 'review.json'
        self.p.save_session(session)
        restored = Project(self.source, BASE / 'templates')
        restored.load_session(session)
        self.assertFalse(restored.normalizations())

    def test_contextual_duplicates_refresh_without_deletion(self):
        self.assertEqual(self.p.duplicate_users(), [])
        first, second, third = self.p.records['USERS']
        second.values = list(first.values)
        group = self.p.duplicate_users()[0]
        self.assertEqual(group['classification'], 'Dados de saída iguais')
        second.values[3] = 'outro@example.com'
        self.assertEqual(self.p.duplicate_users()[0]['differences'], [3])
        self.assertEqual(self.p.duplicate_users()[0]['classification'], 'Dados divergentes')
        second.values[0] = 'Outra Pessoa'
        self.assertEqual(self.p.duplicate_users()[0]['classification'], 'Nomes diferentes')
        third.values[2] = first.values[2]
        self.assertEqual(len(self.p.duplicate_users()[0]['records']), 3)
        self.p.set_value('USERS', [second.row], 2, '99999999999')
        self.assertEqual(len(self.p.duplicate_users()[0]['records']), 2)
        self.assertEqual(len(self.p.analyze().records['USERS']), 3)
        first.values[2] = third.values[2] = ''
        self.assertEqual(self.p.duplicate_users(), [])

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = self.root / 'source.xlsx'
        source_file(self.source)
        self.p = Project(self.source, BASE / 'templates')

    def tearDown(self):
        self.tmp.cleanup()

    def test_mapping_and_normalization(self):
        a = self.p.analyze()
        self.assertEqual({k: len(v) for k, v in a.records.items()}, {'EMPLOYER': 1, 'CUST': 1, 'EXPENSES': 1, 'USERS': 3})
        self.assertEqual(a.records['EMPLOYER'][0].values[1], '01234567000195')
        self.assertEqual(a.records['EMPLOYER'][0].values[13], '01234567')
        self.assertEqual(a.records['EMPLOYER'][0].values[19], 'BRA')
        self.assertEqual(a.records['CUST'][0].values[1], '001')
        self.assertEqual(a.records['USERS'][0].values[6], 'S')
        self.assertIsInstance(a.records['USERS'][0].values[4], datetime)

    def test_mojibake_encoding_fix(self):
        self.p.set_value('CUST', [2], 2, 'DireÃ§Ã£o Executiva (CEO)')
        self.assertEqual(self.p.effective()['CUST'][0].values[2], 'Direção Executiva (CEO)')

    def test_no_automatic_company_or_ambiguous_code(self):
        a = self.p.analyze()
        self.assertEqual(a.records['USERS'][0].values[5], '')
        self.assertEqual(a.records['USERS'][1].values[5], '')
        self.assertTrue(all(r.values[9] == '' for r in a.records['USERS']))
        self.assertTrue(any(i.row == 2 and i.col == 5 for i in a.issues))

    def test_suggestions_are_unique_and_do_not_apply(self):
        a = self.p.analyze()
        self.assertEqual(len(a.suggestions), 2)
        self.assertEqual(a.suggestions[0].proposed, '01234567890@example.com')
        self.assertEqual(a.suggestions[1].proposed, '11144477735@example.com')
        self.assertEqual(a.records['USERS'][0].values[3], 'FINANCEIRO@example.com')

    def test_accept_requires_domain_confirmation(self):
        s = self.p.analyze().suggestions[0]
        with self.assertRaises(ValueError):
            self.p.accept_email(s.row, s.proposed)

    def test_batch_applies_both_duplicates_from_same_snapshot(self):
        self.p.set_value('USERS', [2, 3], 9, '01234567000195')
        self.p.confirm_domain('01234567000195', 'example.com')
        preview = self.p.analyze().suggestions
        self.assertEqual(self.p.accept_all_emails(preview), (2, 0))
        self.assertEqual([r.values[3] for r in self.p.analyze().records['USERS'][:2]], [s.proposed for s in preview])
        self.assertFalse(self.p.analyze().suggestions)

    def test_batch_leaves_unconfirmed_rows_unchanged(self):
        self.p.set_value('USERS', [2], 9, '01234567000195')
        self.p.confirm_domain('01234567000195', 'example.com')
        preview = self.p.analyze().suggestions
        self.assertEqual(self.p.accept_all_emails(preview), (1, 1))
        self.assertEqual(self.p.analyze().records['USERS'][1].values[3], 'financeiro@example.com')

    def test_stale_batch_is_rejected_without_partial_changes(self):
        self.p.set_value('USERS', [2, 3], 9, '01234567000195')
        self.p.confirm_domain('01234567000195', 'example.com')
        preview = self.p.analyze().suggestions
        self.p.set_value('USERS', [3], 0, 'Nome Alterado')
        before = copy.deepcopy(self.p.decisions)
        with self.assertRaises(ValueError):
            self.p.accept_all_emails(preview)
        self.assertEqual(self.p.decisions, before)

    def test_explicit_domain_approval_does_not_assign_company(self):
        preview = self.p.analyze().suggestions
        self.assertEqual(self.p.accept_all_emails(preview, approve_displayed_domains=True), (2, 0))
        self.assertFalse(self.p.analyze().suggestions)
        self.assertTrue(all(r.values[9] == '' for r in self.p.analyze().records['USERS']))
        self.assertFalse(self.p.domains)
        self.assertTrue(any(h[3] == 'Domínio da sugestão aprovado neste lote' for h in self.p.history))

    def test_unit_domain_controls_email_and_explicit_acceptance(self):
        self.p.set_value('USERS', [2, 3], 9, '01234567000195')
        self.p.confirm_domain('01234567000195', '@unidade.example.com')
        s = self.p.analyze().suggestions[0]
        self.assertEqual(s.proposed, '01234567890@unidade.example.com')
        self.p.accept_email(s.row, s.proposed)
        self.assertEqual(self.p.analyze().records['USERS'][0].values[3], s.proposed)

    def test_stale_suggestion_rejected(self):
        self.p.set_value('USERS', [2], 9, '01234567000195')
        self.p.confirm_domain('01234567000195', 'example.com')
        s = self.p.analyze().suggestions[0]
        self.p.confirm_domain('01234567000195', 'another.example.com')
        with self.assertRaises(ValueError):
            self.p.accept_email(s.row, s.proposed)

    def test_invalid_dates_not_guessed(self):
        self.p.set_value('USERS', [2], 4, '13/09/198')
        a = self.p.analyze()
        self.assertEqual(a.records['USERS'][0].values[4], '13/09/198')
        self.assertTrue(any(i.kind == 'USERS' and i.row == 2 and i.col == 4 for i in a.blocking))

    def test_cep_and_phone_normalizations(self):
        self.p.original_values['EMPLOYER:2'][13] = '234567'
        self.p.original_values['EMPLOYER:2'][3] = '(11) 99999-9999'
        proposals = self.p.normalizations()
        categories = {p['category'] for p in proposals}
        self.assertIn('CEP', categories)
        self.assertIn('Telefones', categories)
        cep_prop = next(p for p in proposals if p['category'] == 'CEP')
        self.assertEqual(cep_prop['after'], '00234567')

    def test_cnpj_checksum_validation(self):
        self.p.set_value('EMPLOYER', [2], 1, '01234567000189')  # Invalid CNPJ verifier
        a = self.p.analyze()
        self.assertTrue(any(i.kind == 'EMPLOYER' and i.row == 2 and i.col == 1 and 'CNPJ inválido' in i.message for i in a.issues))

    def test_duplicate_cpf_keeps_both_records(self):
        self.p.set_value('USERS', [3], 2, '01234567890')
        a = self.p.analyze()
        self.assertEqual(len(a.records['USERS']), 3)
        self.assertEqual(len([i for i in a.issues if i.col == 2 and i.kind == 'USERS' and 'repetido' in i.message]), 2)

    def test_draft_export_preserves_template_and_identifiers(self):
        original_hash = file_hash(self.source)
        before = {k: file_hash(v) for k, v in self.p.templates.items()}
        with self.assertRaises(ValueError):
            self.p.export(self.root)
        out = self.p.export(self.root, draft=True)
        self.assertTrue(out.name.startswith('RASCUNHO_'))
        self.assertEqual(file_hash(self.source), original_hash)
        self.assertEqual(before, {k: file_hash(v) for k, v in self.p.templates.items()})
        wb = openpyxl.load_workbook(out / 'DEFAULT_USERS.xlsx')
        self.assertEqual(wb.active['C2'].value, '01234567890')
        self.assertEqual(wb.active['C2'].data_type, 's')
        self.assertIsInstance(wb.active['E2'].value, datetime)
        wb.close()
        report = openpyxl.load_workbook(out / 'REVISAO.xlsx')
        original = list(report['Dados de origem'].values)
        self.assertTrue(any('Conta Contábil' in row and '123.4' in row for row in original))
        report.close()

    def test_session_restores_decisions_and_rejects_different_source(self):
        self.p.set_value('USERS', [2], 9, '01234567000195')
        session = self.root / 'review.json'
        self.p.save_session(session)
        other = Project(self.source, BASE / 'templates')
        other.load_session(session)
        self.assertEqual(other.analyze().records['USERS'][0].values[9], '01234567000195')
        wb = openpyxl.load_workbook(self.source)
        wb['Colaboradores']['A2'] = 'Outro Nome'
        wb.save(self.source)
        wb.close()
        changed = Project(self.source, BASE / 'templates')
        with self.assertRaises(ValueError):
            changed.load_session(session)

    def test_passwords_not_in_report_or_session(self):
        self.p.set_value('USERS', [2], 8, '  test-password  ')
        session = self.root / 'review.json'
        self.p.save_session(session)
        self.assertNotIn('test-password', session.read_text(encoding='utf8'))
        out = self.p.export(self.root, draft=True)
        report = openpyxl.load_workbook(out / 'REVISAO.xlsx')
        self.assertNotIn('test-password', str([[r for r in s.values] for s in report]))
        report.close()

    def test_text_formula_cannot_execute_after_export(self):
        self.p.set_value('USERS', [2], 0, '=1+1')
        out = self.p.export(self.root, draft=True)
        wb = openpyxl.load_workbook(out / 'DEFAULT_USERS.xlsx')
        self.assertEqual(wb.active['A2'].value, '=1+1')
        self.assertEqual(wb.active['A2'].data_type, 's')
        wb.close()

    def test_missing_source_sheet_fails_clearly(self):
        wb = openpyxl.load_workbook(self.source)
        del wb['Colaboradores']
        wb.save(self.source)
        wb.close()
        with self.assertRaisesRegex(ValueError, 'Aba ausente'):
            Project(self.source, BASE / 'templates')

    def test_changed_source_blocks_export(self):
        wb = openpyxl.load_workbook(self.source)
        wb['Colaboradores']['A2'] = 'Novo Nome'
        wb.save(self.source)
        wb.close()
        with self.assertRaisesRegex(ValueError, 'origem foi alterada'):
            self.p.export(self.root, draft=True)

    def test_session_cannot_overwrite_source(self):
        original = file_hash(self.source)
        with self.assertRaises(ValueError):
            self.p.save_session(self.source)
        self.assertEqual(original, file_hash(self.source))

    def test_required_decisions_enable_final_export(self):
        self.p.accept_normalizations(self.p.normalizations())
        self.p.set_value('CUST', [2], 0, 'RAIZ')
        self.p.set_value('EXPENSES', [2], 0, 'DESP001')
        self.p.set_value('EXPENSES', [2], 4, 'VALOR')
        self.p.set_value('USERS', [2], 5, '100')
        self.p.set_value('USERS', [3], 5, '200')
        self.p.set_value('USERS', [4], 5, '300')
        self.p.set_value('USERS', [2, 3, 4], 9, '01234567000195')
        self.p.set_value('USERS', [2], 3, 'pessoa1@example.com')
        self.p.set_value('USERS', [3], 3, 'pessoa2@example.com')
        self.assertFalse(self.p.analyze().blocking)
        self.assertTrue(self.p.export(self.root).name.startswith('CARGAS_'))


if __name__ == '__main__':
    unittest.main()
